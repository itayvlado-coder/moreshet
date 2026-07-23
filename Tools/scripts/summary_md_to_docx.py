#!/usr/bin/env python3
"""Converts an interview summary markdown file (Tools/01-skills/summarize-interview.md format)
into a Hebrew RTL .docx matching the style of Brain/05-Interviews/*/summaries/*.docx examples:
bold 15pt title, bold section headers, plain RTL body paragraphs.

Usage: python3 summary_md_to_docx.py input.md output.docx
"""
import sys
import docx
from docx.shared import Pt


def set_rtl_paragraph(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = pPr.makeelement(docx.oxml.ns.qn("w:bidi"), {docx.oxml.ns.qn("w:val"): "1"})
    pPr.append(bidi)


def add_run(paragraph, text, bold=False):
    run = paragraph.add_run(text)
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rtl = rPr.makeelement(docx.oxml.ns.qn("w:rtl"), {docx.oxml.ns.qn("w:val"): "1"})
    rPr.append(rtl)
    return run


def add_run_with_inline_bold(paragraph, text, base_bold=False, base_size=None):
    """Splits text on **...** markers and renders each segment with the right bold state,
    so literal asterisks never leak into the Word output."""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if part == "":
            continue
        run = add_run(paragraph, part, bold=base_bold or (i % 2 == 1))
        if base_size:
            run.font.size = base_size


def convert(input_path, output_path):
    with open(input_path, encoding="utf-8") as f:
        lines = [l.rstrip("\n") for l in f]

    doc = docx.Document()
    first_line = True

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        if stripped.startswith("# "):
            p = doc.add_paragraph()
            set_rtl_paragraph(p)
            add_run_with_inline_bold(p, stripped[2:], base_bold=True, base_size=Pt(15))
            first_line = False
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            set_rtl_paragraph(p)
            add_run_with_inline_bold(p, stripped[3:], base_bold=True)
        elif stripped.startswith("*(") or stripped.startswith("**הערת"):
            p = doc.add_paragraph()
            set_rtl_paragraph(p)
            add_run_with_inline_bold(p, stripped.strip("*"), base_bold=False)
        else:
            p = doc.add_paragraph()
            set_rtl_paragraph(p)
            add_run_with_inline_bold(p, stripped, base_bold=False)

    doc.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python3 summary_md_to_docx.py input.md output.docx")
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
    print(f"Wrote {sys.argv[2]}")
